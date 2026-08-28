import os
import sqlite3
import argparse
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

def set_cell_bg(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_report(db_path="fridge_data.db", output_docx="Fridge_Diagnostic_Report.docx"):
    if not HAS_DOCX:
        print("Error: python-docx is required. Run 'pip install python-docx'.")
        return

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    PRIMARY = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    GREEN = RGBColor(16, 185, 129)
    ORANGE = RGBColor(249, 115, 22)
    TEXT_GRAY = RGBColor(71, 85, 105)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run('REFRIGERATOR TELEMETRY & DIAGNOSTIC REPORT')
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = BLUE
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f'Generated on {datetime.now().strftime("%Y-%m-%d %H:%M")} | SmartFridge Telemetry System')
    sub_run.font.size = Pt(9.5)
    sub_run.font.color.rgb = TEXT_GRAY
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Query Cycles from SQLite
    cycles = []
    if os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        cur = conn.cursor()
        cur.execute("""
            SELECT 
                strftime('%H:%M', MIN(start_time)),
                strftime('%H:%M', MAX(end_time)),
                MAX(duration_sec),
                avg_power,
                avg_voltage,
                cycle_type
            FROM cycles
            WHERE duration_sec >= 120
            GROUP BY strftime('%H:%M', end_time)
            ORDER BY MAX(end_time) DESC LIMIT 20
        """)
        for r in cur.fetchall():
            m = r[2] // 60
            s = r[2] % 60
            c_type = r[5] if r[5] else ("defrost" if r[3] > 160 else "cooling")
            verdict = "🔥 No Frost Defrost" if c_type == "defrost" else "🟢 Cooling (Compressor)"
            cycles.append((r[0], r[1], f"{m}m {s}s", f"{r[3]} W", f"{r[4]} V", verdict))
        conn.close()

    h1 = doc.add_heading('Recorded Operational Cycles', level=1)
    h1.runs[0].font.color.rgb = PRIMARY

    t = doc.add_table(rows=len(cycles)+1, cols=6)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Start', 'End', 'Duration', 'Avg Power', 'Avg Voltage', 'Mode']
    for j, h_text in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = h_text
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.size = Pt(9)
        set_cell_bg(c, '1E293B')
        set_cell_margins(c, 100, 100, 80, 80)

    for i, row_data in enumerate(cycles, start=1):
        r = t.rows[i]
        for j, val in enumerate(row_data):
            c = r.cells[j]
            c.text = val
            c.paragraphs[0].runs[0].font.size = Pt(8.5)
            if '🟢' in val:
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = GREEN
            elif '🔥' in val:
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = ORANGE
            set_cell_bg(c, 'F8FAFC' if i % 2 == 1 else 'FFFFFF')
            set_cell_margins(c, 80, 80, 80, 80)

    doc.save(output_docx)
    print(f"Report successfully saved to {output_docx}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export Refrigerator Telemetry to Word Document (.docx)")
    parser.add_argument("--db", default="fridge_data.db", help="Path to SQLite database")
    parser.add_argument("--out", default="Fridge_Diagnostic_Report.docx", help="Output .docx file path")
    args = parser.parse_args()
    generate_report(args.db, args.out)
